# -*- coding: utf-8 -*-
"""Generate the BrainHack Donostia project write-up Word template.

The template is based on the BrainHack School Project Guide
(https://school-brainhack.github.io/project_guide/) and the associated
project template repository (https://github.com/school-brainhack/project_template).

Requires the third-party package ``python-docx`` (build-time only, NOT a
dependency of the Jekyll site)::

    python -m pip install python-docx

Usage (default output is ``<repo root>/brainhack_project_template.docx``)::

    python tools/make_bhd_template.py [output.docx]
"""

import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = (
    sys.argv[1]
    if len(sys.argv) > 1
    else os.path.join(REPO_ROOT, "brainhack_project_template.docx")
)

ACCENT = RGBColor(0x2E, 0x54, 0x96)
NAVY = RGBColor(0x1F, 0x38, 0x64)
HINT = RGBColor(0x80, 0x80, 0x80)
DARK = RGBColor(0x40, 0x40, 0x40)

doc = Document()

# ---------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

# ---------------------------------------------------------------- base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def style_heading(name, size, color, bold=True, before=14, after=6):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


style_heading("Heading 1", 17, ACCENT, before=18, after=6)
style_heading("Heading 2", 13.5, ACCENT, before=14, after=4)
style_heading("Heading 3", 11.5, DARK, before=10, after=4)

title_style = doc.styles["Title"]
title_style.font.name = "Calibri"
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = ACCENT
title_style.paragraph_format.space_after = Pt(2)

subtitle_style = doc.styles["Subtitle"]
subtitle_style.font.name = "Calibri"
subtitle_style.font.size = Pt(12)
subtitle_style.font.italic = True
subtitle_style.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


# ---------------------------------------------------------------- helpers
def para(
    text="",
    style=None,
    italic=False,
    bold=False,
    size=None,
    color=None,
    align=None,
    indent=None,
    space_after=None,
):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.italic = italic
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def hint(text):
    return para(text, italic=True, size=10, color=HINT, space_after=10)


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p


def hrule():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E5496")
    borders.append(bottom)
    pPr.append(borders)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, size=10, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color


def make_table(headers, rows, widths, header_fill="D9E2F3"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell(c, htext, bold=True, size=10, color=NAVY)
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            bold = i == 0 and len(headers) == 2
            set_cell(cells[i], val, bold=bold, size=10, italic=(not val))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


# ================================================================ cover
para(
    "BRAINHACK DONOSTIA",
    bold=True,
    size=11,
    color=RGBColor(0x7F, 0x7F, 0x7F),
    space_after=0,
)
doc.add_paragraph("Project Template", style="Title")
para(
    "A write-up template for participants. Based on the BrainHack School "
    "Project Guide (school-brainhack.github.io/project_guide).",
    style="Subtitle",
)
hrule()

para("How to use this template", style="Heading 2")
para(
    "Fill in every section below. Replace the grey hint text with your own "
    "content and delete the hints you do not need.",
    space_after=6,
)
bullet("Write in English (the common language of the school).")
bullet("Keep it concise and structured: this document mirrors a project README.")
bullet(
    "Insert figures, plots and tables in the Results section; give each one a caption."
)
bullet(
    "Aim to use at least 3 open-science tools and 1 new skill learned during the school."
)
bullet(
    "When you are done, save/export the file and add it (with your code) to your GitHub project repository."
)
para("", space_after=4)

# ================================================================ metadata
para("Project information", style="Heading 1")
hint(
    "Complete the table. The summary should be 2-3 sentences that a newcomer can understand."
)
meta_rows = [
    ("Project title", ""),
    ("Author(s)", ""),
    ("Affiliation(s)", ""),
    ("Contact email", ""),
    ("Date", ""),
    ("GitHub repository", ""),
    ("Website / slides", ""),
    ("Tags", ""),
    ("Summary", ""),
]
make_table(["Field", "Your answer"], meta_rows, [4.5, 11.5])
para("", space_after=2)

doc.add_page_break()

# ================================================================ definition
para("1. Project definition", style="Heading 1")

para("Background", style="Heading 2")
hint(
    "Describe the scientific context and motivation. What problem are you addressing, and why does it matter?"
)

para("Objectives", style="Heading 2")
hint(
    "List the concrete question(s) your project aims to answer, or the specific goals you set for yourself."
)

para("Tools", style="Heading 2")
hint(
    "Which 3 open-science tools did you use? Examples: Git, GitHub, containers, Python, BIDS, "
    "Jupyter notebooks, Binder. If you did not use 3, explain why."
)
bullet("Tool 1 —")
bullet("Tool 2 —")
bullet("Tool 3 —")

para("Data", style="Heading 2")
hint(
    "Which dataset(s) did you use? Give the source, an access link and the license. "
    "If the data are private or simulated, state it clearly."
)
para("Data source:", bold=True, space_after=2)
para("License / access:", bold=True, space_after=8)

para("Deliverables", style="Heading 2")
hint(
    "What will you produce at the end? Examples: code, a Jupyter notebook, figures, a report, a website, a presentation."
)
bullet("")
bullet("")
bullet("")

# ================================================================ results
para("2. Results", style="Heading 1")

para("Progress overview", style="Heading 2")
hint(
    "Summarise what was done and the overall status of the project. What worked, what did not?"
)

para("Tools and skills I learned during this project", style="Heading 2")
hint(
    "List the new skill(s), method(s) or technolog(ies) you learned during the school "
    "(e.g. machine learning, connectivity estimation, high-performance computing, DataLad) and how you applied them."
)

para("Results", style="Heading 2")
hint(
    "Present your main findings. Insert your figures, plots or tables below and add a caption for each one."
)
para(
    "[ Figure 1. Caption here. ]",
    italic=True,
    color=HINT,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=14,
)
para(
    "[ Figure 2. Caption here. ]",
    italic=True,
    color=HINT,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=14,
)

# ================================================================ conclusion
para("3. Conclusion and acknowledgements", style="Heading 1")
hint(
    "Wrap up: main take-home message, limitations, next steps, and the people or organisations you want to acknowledge."
)

para("4. References", style="Heading 1")
hint("List the references, datasets and documentation you relied on.")

# ================================================================ appendix
doc.add_page_break()
para("Appendix — Self-assessment checklist (optional)", style="Heading 1")
hint(
    "Rate your own project from 1 to 3: 1 = does not meet expectations, "
    "2 = partially meets expectations, 3 = totally meets expectations."
)
criteria = [
    (
        "Use of open-science best practices",
        "Uses 3 open-science tools learned in week 1, or justifies not using them.",
        "",
    ),
    (
        "Skills and technologies learnt",
        "Uses 1 skill, method or technology learned during the school.",
        "",
    ),
    ("Project relevance", "The project is relevant to brain data analysis.", ""),
    (
        "Clarity",
        "The presentation is easy to follow and supported by convincing material.",
        "",
    ),
    ("Bonus — reproducibility", "Highly reproducible project.", ""),
    ("Bonus — technological achievement", "Noteworthy technical achievement.", ""),
    ("Bonus — presentation", "Exciting presentation.", ""),
    ("Bonus — nice brain picture", "A nice brain picture!", ""),
]
make_table(["Criterion", "What is expected", "Rating (1-3)"], criteria, [5.0, 8.0, 3.0])
para("", space_after=2)

# ================================================================ properties
cp = doc.core_properties
cp.title = "BrainHack Donostia - Project Template"
cp.author = "BrainHack Donostia"
cp.subject = "Project write-up template for BrainHack Donostia participants"

doc.save(OUT)
print("written:", OUT)
